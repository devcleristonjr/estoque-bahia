from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import load_workbook

from app import create_app
from app.extensions import db
from app.models.municipio import Municipio
from app.models.territorio import Territorio


def normalize_header(value: str | None) -> str:
    return (value or "").strip().lower().replace("\n", " ")


def find_column(headers: dict[str, int], candidates: list[str]) -> int | None:
    for candidate in candidates:
        if candidate in headers:
            return headers[candidate]
    return None


def parse_float(value):
    try:
        return float(value) if value not in (None, "") else None
    except (TypeError, ValueError):
        return None


def main() -> int:
    parser = argparse.ArgumentParser(description="Importa territórios e municípios da planilha de referência.")
    parser.add_argument("--file", default="Municipios Bahia.xlsx", help="Caminho para a planilha de referência.")
    args = parser.parse_args()

    workbook_path = Path(args.file)
    if not workbook_path.exists():
        print(f"Planilha não encontrada: {workbook_path.resolve()}")
        print("A aplicação continua funcionando normalmente; a importação pode ser executada depois.")
        return 0

    app = create_app()
    with app.app_context():
        created_territorios, created_municipios = import_workbook(workbook_path)
        print(f"Importação concluída: {created_territorios} territórios, {created_municipios} municípios.")
    return 0


def import_workbook(workbook_path: Path) -> tuple[int, int]:
    workbook = load_workbook(workbook_path, data_only=True)
    worksheet = workbook.active
    headers = {
        normalize_header(cell.value): index
        for index, cell in enumerate(next(worksheet.iter_rows(min_row=1, max_row=1)), start=1)
    }

    municipio_col = find_column(headers, ["municipio", "município", "nome do municipio", "nome do município"])
    territorio_col = find_column(headers, ["territorio", "território", "territorio de identidade", "território de identidade"])
    codigo_ibge_col = find_column(headers, ["codigo ibge", "código ibge", "ibge"])
    latitude_col = find_column(headers, ["latitude", "lat"])
    longitude_col = find_column(headers, ["longitude", "lon", "lng"])
    codigo_territorio_col = find_column(headers, ["codigo territorio", "código território", "codigo", "código"])

    if municipio_col is None or territorio_col is None:
        raise RuntimeError("A planilha precisa conter ao menos as colunas de município e território.")

    created_territorios = 0
    created_municipios = 0

    for row in worksheet.iter_rows(min_row=2, values_only=True):
        municipio_nome = row[municipio_col - 1] if len(row) >= municipio_col else None
        territorio_nome = row[territorio_col - 1] if len(row) >= territorio_col else None
        if not municipio_nome or not territorio_nome:
            continue

        territorio_nome = str(territorio_nome).strip()
        municipio_nome = str(municipio_nome).strip()
        territorio = Territorio.query.filter_by(nome=territorio_nome).first()
        if territorio is None:
            territorio = Territorio(
                nome=territorio_nome,
                codigo=str(row[codigo_territorio_col - 1]).strip() if codigo_territorio_col and row[codigo_territorio_col - 1] else None,
                ativo=True,
            )
            db.session.add(territorio)
            db.session.flush()
            created_territorios += 1

        codigo_ibge = str(row[codigo_ibge_col - 1]).strip() if codigo_ibge_col and row[codigo_ibge_col - 1] else None
        municipio = Municipio.query.filter(
            Municipio.nome == municipio_nome,
            Municipio.territorio_id == territorio.id,
        ).first()
        if municipio is None:
            municipio = Municipio(
                nome=municipio_nome,
                territorio=territorio,
                codigo_ibge=codigo_ibge,
                latitude=parse_float(row[latitude_col - 1]) if latitude_col else None,
                longitude=parse_float(row[longitude_col - 1]) if longitude_col else None,
                ativo=True,
            )
            db.session.add(municipio)
            created_municipios += 1
        else:
            municipio.codigo_ibge = municipio.codigo_ibge or codigo_ibge
            if not municipio.latitude and latitude_col:
                municipio.latitude = parse_float(row[latitude_col - 1])
            if not municipio.longitude and longitude_col:
                municipio.longitude = parse_float(row[longitude_col - 1])

    db.session.commit()
    return created_territorios, created_municipios


if __name__ == "__main__":
    raise SystemExit(main())
