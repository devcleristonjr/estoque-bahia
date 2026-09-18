from __future__ import annotations

import os
from pathlib import Path

from dotenv import load_dotenv
from flask import Flask, render_template
from openpyxl import load_workbook

from app.extensions import csrf, db, login_manager, migrate
from app.models.fechamento_diario_estoque import FechamentoDiarioEstoque
from app.models.municipio import Municipio
from app.models.usuario import Usuario
from app.routes.administracao import municipios_bp, territorios_bp, usuarios_bp
from app.routes.api import api_bp
from app.routes.coleta import coleta_bp
from app.routes.coleta_public import coleta_public_bp
from app.routes.auth import auth_bp
from app.routes.dashboard import dashboard_bp
from app.routes.estoques import estoques_bp
from app.routes.files import files_bp
from app.routes.materiais import materiais_bp
from app.commands import register_commands
from app.timezone import formatar_datahora_bahia
from config import get_config


load_dotenv()


@login_manager.user_loader
def load_user(user_id: str) -> Usuario | None:
    if not user_id:
        return None
    return db.session.get(Usuario, int(user_id))


def _ensure_reference_municipal_data() -> None:
    if Municipio.query.count() > 0:
        return

    workbook_path = Path(__file__).resolve().parent.parent / "Municipios Bahia.xlsx"
    if not workbook_path.exists():
        return

    workbook = load_workbook(workbook_path, data_only=True)
    worksheet = workbook.active
    headers = {
        str(cell.value).strip().lower().replace("\n", " "): index
        for index, cell in enumerate(next(worksheet.iter_rows(min_row=1, max_row=1)), start=1)
    }

    def find_column(candidates):
        for candidate in candidates:
            if candidate in headers:
                return headers[candidate]
        return None

    municipio_col = find_column(["municipio", "município", "nome do municipio", "nome do município"])
    territorio_col = find_column(["territorio", "território", "territorio de identidade", "território de identidade"])
    if municipio_col is None or territorio_col is None:
        return

    from app.models.territorio import Territorio

    for row in worksheet.iter_rows(min_row=2, values_only=True):
        municipio_nome = row[municipio_col - 1] if len(row) >= municipio_col else None
        territorio_nome = row[territorio_col - 1] if len(row) >= territorio_col else None
        if not municipio_nome or not territorio_nome:
            continue

        territorio = Territorio.query.filter_by(nome=str(territorio_nome).strip()).first()
        if territorio is None:
            territorio = Territorio(nome=str(territorio_nome).strip(), ativo=True)
            db.session.add(territorio)
            db.session.flush()

        if not Municipio.query.filter_by(nome=str(municipio_nome).strip(), territorio_id=territorio.id).first():
            db.session.add(Municipio(nome=str(municipio_nome).strip(), territorio=territorio, ativo=True))

    db.session.commit()


def create_app(config_object: type | None = None) -> Flask:
    app = Flask(__name__)
    app.config.from_object(config_object or get_config())

    if not app.config.get("SECRET_KEY"):
        raise RuntimeError("SECRET_KEY is required. Create a .env file based on .env.example.")
    if not app.config.get("SQLALCHEMY_DATABASE_URI"):
        raise RuntimeError("DATABASE_URL is required. Create a PostgreSQL database and set it in .env.")

    upload_folder = Path(app.config["UPLOAD_FOLDER"])
    upload_folder.mkdir(parents=True, exist_ok=True)
    (upload_folder / "estoque").mkdir(parents=True, exist_ok=True)

    db.init_app(app)
    migrate.init_app(app, db)
    login_manager.init_app(app)
    csrf.init_app(app)

    with app.app_context():
        db.create_all()
        _ensure_reference_municipal_data()

    login_manager.login_view = "auth.login"
    login_manager.login_message = "Faça login para continuar."
    login_manager.login_message_category = "warning"

    @app.template_filter("datahora_bahia")
    def datahora_bahia_filter(value, pattern: str = "%d/%m/%Y %H:%M"):
        return formatar_datahora_bahia(value, pattern)

    app.register_blueprint(auth_bp)
    app.register_blueprint(dashboard_bp)
    app.register_blueprint(estoques_bp)
    app.register_blueprint(coleta_bp)
    app.register_blueprint(coleta_public_bp)
    app.register_blueprint(materiais_bp)
    app.register_blueprint(usuarios_bp)
    app.register_blueprint(territorios_bp)
    app.register_blueprint(municipios_bp)
    app.register_blueprint(files_bp)
    app.register_blueprint(api_bp)

    register_commands(app)

    @app.errorhandler(403)
    def forbidden(_error):
        return render_template("errors/403.html"), 403

    @app.errorhandler(404)
    def not_found(_error):
        return render_template("errors/404.html"), 404

    return app
